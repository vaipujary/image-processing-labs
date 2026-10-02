"""Build the editable Lab 4 report from saved results and measured figures.

Run using the Codex bundled document Python or an environment with python-docx
and lxml. Text, tables, and display equations remain native Word objects.
"""
from copy import deepcopy
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

HERE=Path(__file__).resolve().parent
RESULTS=HERE/'results'
OUTPUT=HERE/'output/docx/Lab_4_Report.docx'


def element(name,children=(),**attrs):
    node=OxmlElement(name)
    for key,value in attrs.items(): node.set(qn(key),str(value))
    for child in children: node.append(deepcopy(child))
    return node


def text(value):
    node=element('m:t'); node.text=value; node.set(qn('xml:space'),'preserve')
    return element('m:r',[node])


def script(base,sub=None,sup=None):
    nodes=[element('m:e',[text(base)] if isinstance(base,str) else base)]
    if sub is not None: nodes.append(element('m:sub',[text(sub)]))
    if sup is not None: nodes.append(element('m:sup',[text(sup)]))
    return element('m:sSubSup' if sub is not None and sup is not None else 'm:sSub' if sub is not None else 'm:sSup',nodes)


def frac(num,den):
    return element('m:f',[element('m:num',num),element('m:den',den)])


def equation(doc,nodes):
    p=doc.add_paragraph()
    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(7)
    p._p.append(element('m:oMathPara',[
        element('m:oMathParaPr',[element('m:jc',**{'m:val':'center'})]),element('m:oMath',nodes)]))
    return p


def paragraph(doc,value,style=None):
    return doc.add_paragraph(value,style)


def hyperlink(paragraph,label,url):
    relation=paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True)
    node=element('w:hyperlink',**{'r:id':relation})
    value=element('w:t');value.text=label
    node.append(element('w:r',[element('w:rPr',[element('w:u',**{'w:val':'single'})]),value]))
    paragraph._p.append(node)


def heading(doc,value):
    return doc.add_heading(value,level=1)


def figure(doc,name,width,caption=None):
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.keep_with_next=bool(caption)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=p.add_run();run.add_picture(str(RESULTS/name),width=Inches(width))
    docpr=run._r.xpath('.//wp:docPr')
    if docpr: docpr[0].set('descr',caption or name.replace('_',' '))
    if caption: paragraph(doc,caption,'Caption')


def page(doc):
    doc.add_page_break()


def configure(doc):
    for border in doc.styles.element.xpath('.//w:pBdr'): border.getparent().remove(border)
    section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=Inches(.65);section.bottom_margin=Inches(.65)
    section.left_margin=Inches(.75);section.right_margin=Inches(.75)
    section.footer_distance=Inches(.25)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption','Quote']:
        s=doc.styles[name]
        s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
        s.font.size=Pt(10.5)
        s.paragraph_format.space_after=Pt(6)
    normal=doc.styles['Normal']
    normal.paragraph_format.line_spacing=1.06
    for name,size in [('Title',23),('Subtitle',13),('Heading 1',13),('Heading 2',11),('Caption',9)]:
        doc.styles[name].font.size=Pt(size)
    doc.styles['Title'].font.bold=True
    doc.styles['Heading 1'].font.bold=True
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(9)
    doc.styles['Heading 1'].paragraph_format.space_after=Pt(5)
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(7)
    doc.styles['Heading 2'].paragraph_format.space_after=Pt(3)
    doc.styles['Caption'].font.italic=False
    doc.styles['Caption'].font.bold=False
    doc.styles['Quote'].font.italic=False
    doc.styles['Quote'].font.size=Pt(9.5)
    doc.styles['Quote'].paragraph_format.left_indent=Inches(.18)
    doc.styles['Quote'].paragraph_format.right_indent=Inches(.1)
    footer=section.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=footer.add_run('Lab 4  |  ');r.font.size=Pt(8)
    field=element('w:fldSimple',**{'w:instr':'PAGE'})
    value=element('w:t');value.text='1'
    field.append(element('w:r',[element('w:rPr',[element('w:sz',**{'w:val':'16'})]),value]))
    footer._p.append(field)
    doc.core_properties.title='Lab 4 MAP Reconstruction with Non Gaussian Priors'
    doc.core_properties.author='Vaidehi Pujary'
    doc.core_properties.subject='ECE 60141 Lab 4 D1 through D11'


def results_table(doc,s):
    headers=['Run','Prior','T','σx','Final true cost','RMSE','Time s']
    table=doc.add_table(rows=1,cols=len(headers))
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    widths=[.42,1.32,.42,.60,1.52,.91,.69]
    for column,width in zip(table.columns,widths): column.width=Inches(width)
    for cell,value in zip(table.rows[0].cells,headers): cell.text=value
    for key,r in s['runs'].items():
        cells=table.add_row().cells
        values=[key,r['prior'],f"{r['T']:g}",f"{r['sigma_x']:.3f}",f"{r['final_cost']:,.3f}",f"{r['rmse']:.8f}",f"{r['seconds']:.2f}"]
        for c,v in zip(cells,values):c.text=v
    for ri,row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(element('w:cantSplit'))
        if ri==0: row._tr.get_or_add_trPr().append(element('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc=c._tc.get_or_add_tcPr()
            tc.append(element('w:shd',**{'w:fill':'D9E5F2' if ri==0 else 'FFFFFF'}))
            tc.append(element('w:tcMar',[element('w:top',**{'w:w':'80','w:type':'dxa'}),element('w:bottom',**{'w:w':'80','w:type':'dxa'})]))
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0)
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci<2 else WD_ALIGN_PARAGRAPH.RIGHT
                for r in p.runs:r.font.size=Pt(8.5);r.bold=ri==0
    borders=element('w:tblBorders',[element('w:'+n,**{'w:val':'single','w:sz':'4','w:color':'D9D9D9'}) for n in ['top','left','bottom','right','insideH','insideV']])
    table._tbl.tblPr.append(borders)
    return table


def main():
    s=json.loads((RESULTS/'summary.json').read_text())
    v=json.loads((RESULTS/'verification.json').read_text())
    r=s['runs']; improvement=100*(1-r['R2']['rmse']/r['R1']['rmse'])
    doc=Document();configure(doc)
    doc.add_paragraph('Lab 4',style='Title')
    doc.add_paragraph('MAP Reconstruction with Non Gaussian Priors',style='Subtitle')
    paragraph(doc,'ECE 60141 Foundations of Computational Imaging')
    paragraph(doc,'Vaidehi Pujary  •  October 1 2026  •  Random seed 0')
    hyperlink(paragraph(doc,'GitHub repository: '),'https://github.com/vaipujary/image-processing-labs','https://github.com/vaipujary/image-processing-labs')
    paragraph(doc,f'The QGGMRF reconstruction preserves the parrot’s beak edge more clearly than the Gaussian comparison and reduces full-image RMSE from {r["R1"]["rmse"]:.5f} to {r["R2"]["rmse"]:.5f} ({improvement:.1f}%). The true MAP cost decreases at every update in all six 50-iteration runs.')
    heading(doc,'Experimental setup')
    paragraph(doc,'The 512 by 768 image kodim23.pgm is scaled to [0, 1]. Levin kernel 1 defines circular convolution, and one Gaussian noise image with σw = 0.02 is shared by all runs. The observation is identical to the saved Lab 3 deconvolution observation. Neighbor pairs use only pixels inside the image, with axial weights 1/6 and diagonal weights 1/12.')
    paragraph(doc,'Every run uses float64 on CPU, p = 2, q = 1.2, ω = 1, and exactly 50 updates from x = y. The true gradient is computed with torch.autograd; a fresh quadratic surrogate supplies the directional step size. No image or measurement is clipped in the calculations. All reconstruction images use the same display range [0, 1].')
    heading(doc,'D1 Potential and influence functions')
    figure(doc,'D1_potential_influence.png',6.85,'Figure 1. Potential and influence over −0.3 ≤ Δ ≤ 0.3, with σx = 0.2, p = 2, q = 1.2 and T = 1. The blue comparison is the exact quadratic potential of Eq. (2).')
    paragraph(doc,'The quadratic prior has a linearly increasing influence, so it strongly pulls together the unequal pixels on opposite sides of an edge. The QGGMRF has a smaller, sublinearly growing influence at large differences, so it permits a sharper jump while retaining strong smoothing near zero.')
    paragraph(doc,'For q = 1.2, the QGGMRF influence remains unbounded but grows asymptotically as the 0.2 power of |Δ|. This slower growth explains its edge-preserving behavior.')

    page(doc)
    heading(doc,'D2 Derivation of the zero difference limit')
    paragraph(doc,'Set p = 2 in Eq. (14), so the prefactor |Δ′|ᵖ⁻² is one. For nonzero Δ′ define z as follows; because q < 2, z tends to infinity as Δ′ tends to zero.')
    equation(doc,[text('z = '),script([text('|'),frac([text('Δ′')],[text('T'),script('σ',sub='x')]),text('|')],sup='q−2'),text(' → ∞')])
    prefactor=frac([script('b',sub='s,r')],[text('2'),script('σ',sub='x',sup='2')])
    equation(doc,[script('b̃',sub='s,r'),text(' = '),prefactor,
                  frac([text('z(q/2 + z)')],[script('(1 + z)',sup='2')]),text(' = '),prefactor,
                  frac([text('1 + (q/2)'),script('z',sup='−1')],[text('1 + 2'),script('z',sup='−1'),text(' + '),script('z',sup='−2')])])
    paragraph(doc,'The second equality divides numerator and denominator by z². In particular, the leading powers of Δ′ cancel explicitly:')
    equation(doc,[frac([script('z',sup='2')],[script('z',sup='2')]),text(' = '),
                  frac([script('|Δ′|',sup='2(q−2)'),script([text('(T'),script('σ',sub='x'),text(')')],sup='−2(q−2)')],
                       [script('|Δ′|',sup='2(q−2)'),script([text('(T'),script('σ',sub='x'),text(')')],sup='−2(q−2)')]),text(' = 1')])
    paragraph(doc,'As z tends to infinity, both inverse powers vanish. Therefore the finite limiting coefficient is Eq. (15):')
    equation(doc,[element('m:limLow',[element('m:e',[text('lim')]),element('m:lim',[text('Δ′→0')])]),text(' '),script('b̃',sub='s,r'),text(' = '),prefactor,text('.')])
    paragraph(doc,'The implementation uses this limit at |Δ′| ≤ 10⁻⁸. It evaluates the other branch using the reciprocal variable u, whose exponent is p − q > 0, after flooring the power argument. Thus neither evaluated branch contains an infinite negative power.')
    heading(doc,'D3 Touching quadratic surrogates')
    figure(doc,'D3_surrogates.png',6.85,'Figure 2. The potential and its shifted symmetric quadratic bounds, using the D1 parameters; dots mark Δ = ±0.02 and ±0.15.')
    paragraph(doc,'Both surrogates touch the potential at their two marked symmetric points and lie above it elsewhere; the numerical grid check also confirms the bound to floating-point tolerance.')

    page(doc)
    heading(doc,'D4 AI specification and verification')
    paragraph(doc,'The actual request given to the assistant was:')
    paragraph(doc,'“can you do this whole lab: https://cabouman.github.io/grad_labs/labs-60141/lab4/index.html? follow all the instructions on the webpage. for the report, create a word document and put the contents in so that I can also make updates to it after you\'re done. the textbook for the course is included in this directory if you want to refer to it. the lab3 files are also in this directory”','Quote')
    paragraph(doc,'The assistant translated the linked instructions into this implementation specification; this is not a separate student-written prompt:')
    paragraph(doc,'“Implement the eight required signatures in a functions-only lab4.py. Reuse the four permitted Lab 3 functions. Use float64, no pixel loops, and unconstrained images. Evaluate Eqs. (1), (3), and (4), using the prescribed (8, H, W) differences and validity mask in NW, N, NE, W, E, SW, S, SE order. Use 1/6 axial and 1/12 diagonal weights and exclude pairs outside the image. Floor the fractional-power argument at 10⁻⁸; use Eq. (15) for surrogate weights at or below that threshold and the equivalent reciprocal form of Eq. (14) otherwise. Compute the true gradient with autograd and check Eq. (16). Use α = (gᵀg)/(gᵀHg) from Eqs. (17)–(18), without an extra half in the Hessian double sum. Start at y; perform 50 updates x ← x − ωαg and return 51 true costs. Share seed-0 noise across R1–R6, with σw = 0.02, p = 2, q = 1.2 and ω = 1. Use every listed (T, σx) pair, generate D1–D11, and report unmodified RMSEs and timed results.”','Quote')
    paragraph(doc,'The first test command failed with “lab4 is not a package” before any numerical tests ran. Adding __init__.py resolved the module/package name collision. The first numerical implementation then passed six test groups: potential/influence, neighbor masks and independent pair counting, analytic-versus-autograd gradient, zero-limit weights and Hessian curvature, scalar majorization, and descent including a stationary input.')
    paragraph(doc,f'Independent verification uses FFT convolution and four unordered-pair slice sums. The maximum gradient discrepancy is {v["autograd_vs_equation16_max_abs_error"]:.2e}; the largest final-cost discrepancy is {max(item["cost_abs_error"] for item in v["runs"].values()):.2e}. No numerical formula correction was needed. The full instruction and correction record is AI_SPECIFICATION.md.')
    heading(doc,'D5 True cost decreases at every step')
    figure(doc,'D5_true_cost.png',6.4)
    paragraph(doc,f'Figure 3. R2 decreases strictly at all 50 updates, from {r["R2"]["cost_history"][0]:,.3f} to {r["R2"]["final_cost"]:,.3f}; even the smallest decrease is {abs(r["R2"]["largest_cost_change"]):.6f}. These are true costs, not surrogate values.')

    page(doc)
    heading(doc,'D6 Reconstructions and common crop')
    paragraph(doc,'The four full images use grayscale [0, 1]. Orange boxes locate the common 128 by 128 crop: rows 160–287 and columns 384–511, using zero-based indices, or [160:288, 384:512] in Python.')
    figure(doc,'D6_full_images.png',6.85,'Figure 4. Original image, shared blurred/noisy observation, R1 (T = 100, σx = 0.036), and R2 (T = 1, σx = 0.020), after 50 updates.')
    figure(doc,'D6_crops.png',6.85,'Figure 5. The same four images in the fixed crop, shown with the same [0, 1] grayscale. The crop includes the bright beak and its dark right-hand edge.')
    paragraph(doc,f'Observation RMSE is {s["input_rmse"]:.8f}. R1 has RMSE {r["R1"]["rmse"]:.8f}; R2 has RMSE {r["R2"]["rmse"]:.8f}. The QGGMRF preserves more beak-edge contrast and fine face structure, with some residual texture and ringing. Only display rendering saturates values outside [0, 1].')

    page(doc)
    heading(doc,'D7 Edge profile')
    paragraph(doc,'The horizontal profile uses row 232 and columns 384–511 of the same crop. The right panel enlarges columns 438–460 around the bright-to-dark beak transition; both panels contain the original, R1, and R2.')
    figure(doc,'D7_edge_profile.png',6.85,'Figure 6. Pixel values in [0, 1] image units, with unclipped estimates. The QGGMRF follows the transition more closely, although neither 50-step reconstruction fully recovers the original edge.')
    heading(doc,'D8 Prior scale study')
    figure(doc,'D8_scale_study.png',6.6,'Figure 7. Half, center, and double σx from left to right. Top: QGGMRF R3, R2, R4. Bottom: Gaussian comparison R5, R1, R6. Labels give full-image RMSE; every crop uses [0, 1] grayscale.')

    page(doc)
    heading(doc,'D9 Sum of the final surrogate weights')
    paragraph(doc,'The weights are recomputed at the final R2 estimate x⁽⁵⁰⁾ and summed over all eight neighbor planes. The displayed range is 0 to 1250; 1250 = 1/(2 × 0.020²) is the maximum interior sum at equal neighbors.')
    figure(doc,'D9_weight_sum.png',6.75,'Figure 8. Summed weights at R2 iteration 50. Dark contours indicate reduced coupling across edges; bright areas indicate stronger smoothing of locally similar neighbors.')
    paragraph(doc,f'The observed range is {s["figures"]["weight_actual_range"][0]:.2f} to {s["figures"]["weight_actual_range"][1]:.2f}. Fine facial and feather structure appears in the weights. Lower sums along the image boundary also reflect missing neighbors, not only image contrast.')
    heading(doc,'D10 Measured results for all six runs')
    results_table(doc,s)
    paragraph(doc,'Table 1. Final true cost, full-image RMSE in [0, 1] units, and elapsed reconstruction time for exactly 50 updates. Timings include cost recording, exclude image loading and plotting, and follow a short warmup.')
    paragraph(doc,f'Execution used Python {s["python"]}, PyTorch {s["torch"]}, CPU and {s["threads"]} threads. Times vary with machine load. All 300 cost changes were negative. The costs use different prior parameters and cannot be ranked across rows as a measure of reconstruction quality; RMSE uses the same ground truth in every row.')
    paragraph(doc,'The lab calls T = 100 the Gaussian MRF. The implementation retains the prescribed QGGMRF formula at T = 100, which approximates the quadratic regime; it does not replace it with an exactly quadratic cost. The D1 blue curve is the exact quadratic comparison.')

    page(doc)
    heading(doc,'D11 Answers to the seven questions')
    questions=[
        ('1 Edge and flat region behavior','R2 retains a steeper transition and stronger local contrast along the beak edge than R1 (Figures 5 and 6); R1 blends the two sides more strongly. Both reduce noise in relatively flat regions, while R2 retains more small-scale structure. Figure 1 explains this: the QGGMRF influence grows more slowly for large pixel differences, so an edge is penalized less strongly.'),
        ('2 RMSE and visual quality',f'Yes. R2 has the sharper-looking central reconstruction and the lower RMSE, {r["R2"]["rmse"]:.5f} versus {r["R1"]["rmse"]:.5f}, a {improvement:.1f}% reduction. RMSE measures global squared error rather than perceptual edge quality, so this agreement does not make it a complete measure of sharpness or ringing.'),
        ('3 What the prior scale controls',f'Smaller σx strengthens regularization. For QGGMRF, RMSE is {r["R3"]["rmse"]:.5f}, {r["R2"]["rmse"]:.5f}, and {r["R4"]["rmse"]:.5f} at half, center, and double scale; for the Gaussian comparison it is {r["R5"]["rmse"]:.5f}, {r["R1"]["rmse"]:.5f}, and {r["R6"]["rmse"]:.5f}. At the converged optimum, σx → 0 forces near-constant images on the connected grid; σx → ∞ removes the prior and approaches unregularized least squares, amplifying noise in weakly observed blur modes.'),
        ('4 Why the QGGMRF can use a smaller scale','At small differences, the limiting curvature is 1/σx², so R2 can smooth flat regions more strongly with σx = 0.020 than R1 with 0.036. Across an edge, the QGGMRF influence and surrogate weights decrease relative to the quadratic law, allowing that stronger local smoothing without the same loss of edge contrast. The model adapts its coupling to each pair’s difference.'),
        ('5 Why p equals two','In the allowed range 1 ≤ q < p ≤ 2, choosing p = 2 gives the finite zero-difference curvature ρ″(0) = 1/σx² and hence the finite coefficient in D2. With p < 2 the curvature diverges near zero, so this surrogate cannot use a finite limiting weight. The large-difference exponent q remains a modeling choice within 1 ≤ q < 2.'),
        ('6 Why the true cost is monotone','Yes: all 50 R2 changes are negative (D5). The exact surrogate line step decreases Q, giving the middle inequality in Eq. (8); the upper bound and touching conditions then transfer that decrease to the true cost. Weights that are too small can break the upper bound, so a reduced surrogate cost would no longer guarantee a reduced true cost.'),
        ('7 Structure in the weight map','Figure 8 contains dark outlines of the beaks, eyes, facial markings, and feathers, with higher weights in smoother regions. Large neighbor differences reduce ρ′(Δ)/(2Δ), weakening the smoothing across those structures; small differences give larger weights. The summed map combines all directions, so a visible edge can still have substantial weight from neighbors along the edge.')]
    for title,answer in questions:
        doc.add_heading(title,level=2);paragraph(doc,answer)
    doc.add_heading('References',level=2)
    paragraph(doc,'[1] C. A. Bouman, ECE 60141 Lab 4, MAP Reconstruction with Non-Gaussian Priors, live instructions accessed October 1, 2026. https://cabouman.github.io/grad_labs/labs-60141/lab4/index.html')
    paragraph(doc,'[2] C. A. Bouman, Foundations of Computational Imaging: A Model-Based Approach, SIAM, 2022. Chapter 6, §§6.3–6.4, pp. 89–94; Chapter 8, §§8.1–8.4, pp. 109–119. The lab uses p for the near-zero exponent and q for the large-difference exponent; the book reverses these names.')
    OUTPUT.parent.mkdir(parents=True,exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__=='__main__':
    main()
